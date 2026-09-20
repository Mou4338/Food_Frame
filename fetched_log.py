"""Cross-run 'already fetched' registry, stored as its own Excel file --
separate from BOTH the input menu Excel (which only tracks one run) and the
SQLite database (which isn't something a non-technical worker can casually
open and read).

Point a brand-new input menu file at the same dish next month, and if it's
already in this log, the pipeline skips searching/scoring/AI-checking it
again and just reuses the Drive link already on file. This is checked by
*normalized* name (see name_cleaner.normalize_name), so "Chicken Tandoori",
"chicken  tandoori", and "Chicken Tandoori [Half]" all count as the same
already-fetched dish.

The file is created automatically on first use -- nothing to set up beyond
pointing FETCHED_LOG_PATH at where you want it (default ./fetched_items.xlsx).
"""
import os
from datetime import datetime

import openpyxl

from name_cleaner import normalize_name

HEADERS = ["food_name", "normalized_name", "drive_link", "source", "fetched_at"]


def load_fetched(path: str) -> dict:
    """Returns {normalized_name: {"food_name", "drive_link", "source", "fetched_at"}}.
    Returns {} if the log doesn't exist yet (first run)."""
    if not path or not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        food_name, normalized, drive_link, source, fetched_at = (list(row) + [None] * 5)[:5]
        key = normalized or normalize_name(food_name)
        if key:
            out[key] = {"food_name": food_name, "drive_link": drive_link, "source": source, "fetched_at": fetched_at}
    wb.close()
    return out


def record_fetched(path: str, food_name: str, drive_link: str, source: str = ""):
    """Appends one row. Creates the file (with headers) the first time it's called."""
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Fetched"
        ws.append(HEADERS)
    ws.append([
        food_name, normalize_name(food_name), drive_link or "", source or "",
        datetime.now().isoformat(timespec="seconds"),
    ])
    wb.save(path)
