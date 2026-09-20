"""Reads food item rows from Excel and writes results back.

Handles real-world sheets where the header isn't necessarily row 1, and
maps every duplicate/near-duplicate dish name to the SAME processing job
(via name_cleaner.normalize_name) so 502 menu rows with 211 unique dishes
only get searched/scored once -- the result is then written back to every
matching row.
"""
import openpyxl
from name_cleaner import normalize_name


class SheetManager:
    def __init__(self, excel_path: str, sheet_name: str, food_name_column: str,
                 category_column: str | None = None, max_header_scan: int = 10):
        self.excel_path = excel_path
        self.wb = openpyxl.load_workbook(excel_path)
        self.ws = self.wb[sheet_name] if sheet_name in self.wb.sheetnames else self.wb.active

        self.header_row = self._find_header_row(food_name_column, max_header_scan)
        self.headers = [c.value for c in self.ws[self.header_row]]
        self.food_col_idx = self.headers.index(food_name_column) + 1
        self.category_col_idx = (
            self.headers.index(category_column) + 1
            if category_column and category_column in self.headers else None
        )

        self.status_col_idx = self._get_or_create("Status")
        self.link_col_idx = self._get_or_create_link_column()
        self.notes_col_idx = self._get_or_create("Notes")

    def _find_header_row(self, food_name_column: str, max_scan: int) -> int:
        for row in range(1, max_scan + 1):
            values = [c.value for c in self.ws[row]]
            if any(isinstance(v, str) and v.strip() == food_name_column for v in values):
                return row
        raise ValueError(
            f"Could not find a column named '{food_name_column}' in the first {max_scan} rows of the sheet."
        )

    def _get_or_create(self, header_name):
        if header_name in self.headers:
            return self.headers.index(header_name) + 1
        new_idx = len(self.headers) + 1
        self.ws.cell(row=self.header_row, column=new_idx, value=header_name)
        self.headers.append(header_name)
        return new_idx

    def _get_or_create_link_column(self):
        for i, h in enumerate(self.headers):
            if isinstance(h, str) and h.strip().lower() in ("google drive link", "drive link"):
                return i + 1
        return self._get_or_create("Google drive link")

    def rows_to_process(self):
        """Yields (row_number, food_item_name, category) for rows not yet marked Done."""
        for row in range(self.header_row + 1, self.ws.max_row + 1):
            name = self.ws.cell(row=row, column=self.food_col_idx).value
            status = self.ws.cell(row=row, column=self.status_col_idx).value
            if name and str(name).strip() and status != "Done":
                category = (
                    self.ws.cell(row=row, column=self.category_col_idx).value
                    if self.category_col_idx else None
                )
                yield row, str(name).strip(), category

    def group_by_normalized_name(self):
        """Groups all not-yet-done rows by normalized dish name, so each
        distinct dish is processed exactly once. Returns
        {normalized_name: {"display_name": str, "category": str, "rows": [row_num, ...]}}"""
        groups = {}
        for row_num, name, category in self.rows_to_process():
            key = normalize_name(name)
            g = groups.setdefault(key, {"display_name": name, "category": category, "rows": []})
            g["rows"].append(row_num)
        return groups

    def mark_done(self, row: int, drive_link: str, note: str = ""):
        self.ws.cell(row=row, column=self.status_col_idx, value="Done")
        self.ws.cell(row=row, column=self.link_col_idx, value=drive_link)
        if note:
            self.ws.cell(row=row, column=self.notes_col_idx, value=note[:250])

    def mark_pending_review(self, row: int, note: str = ""):
        self.ws.cell(row=row, column=self.status_col_idx, value="Pending Review")
        if note:
            self.ws.cell(row=row, column=self.notes_col_idx, value=note[:250])

    def mark_failed(self, row: int, reason: str):
        self.ws.cell(row=row, column=self.status_col_idx, value="Failed")
        self.ws.cell(row=row, column=self.notes_col_idx, value=reason[:250])

    def save(self, path: str | None = None):
        self.wb.save(path or self.excel_path)
